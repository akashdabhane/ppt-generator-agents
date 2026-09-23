import pytest
from pptx import Presentation as PPTXPresentation
from pptx.util import Emu

from app.presentation.layout_engine import LayoutEngine
from app.presentation.renderer import PresentationRenderer
from app.presentation.themes import THEMES
from app.schemas.presentation_spec import PresentationSpec

LONG = "detailed supporting evidence taken from the quarterly filing and board minutes " * 3
CITE = [{"document_name": "Q3_report.pdf", "page": 4, "section": "Revenue", "excerpt": "..."}]


def _heavy_spec() -> PresentationSpec:
    return PresentationSpec.model_validate({
        "title": "Deck",
        "slides": [
            {"type": "title", "title": "A very long presentation title " * 6, "subtitle": "Subtitle " * 40},
            {"type": "section", "title": "Section", "subtitle": "Overview"},
            {"type": "bullet", "title": "Findings " * 10, "bullets": [f"{i}. {LONG}" for i in range(14)], "citations": CITE},
            {"type": "two_column", "title": "Compare", "left_title": "Challenges " * 8,
             "left_content": [LONG] * 9, "right_title": "Solutions", "right_content": ["Short"], "citations": CITE},
            {"type": "table", "title": "Data", "columns": ["A", "B", "C"],
             "rows": [[f"r{i}", LONG, "x"] for i in range(12)], "citations": CITE},
            {"type": "chart", "title": "Chart", "chart_type": "bar", "labels": ["Q1", "Q2", "Q3"],
             "values": [1.0, 2.0], "citations": CITE},
            {"type": "chart", "title": "Empty chart", "chart_type": "pie", "labels": [], "values": []},
            {"type": "quote", "title": "Quote", "quote": LONG * 4, "author": "CFO " * 30, "citations": CITE},
            {"type": "summary", "title": "Summary", "key_takeaways": [LONG] * 10, "citations": CITE},
        ],
    })


def _texts(prs):
    return [p.text for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs]


@pytest.mark.parametrize("theme", list(THEMES))
def test_heavy_deck_stays_inside_slide_bounds(tmp_path, theme):
    path = PresentationRenderer(theme).render(_heavy_spec(), str(tmp_path / "deck.pptx"))
    prs = PPTXPresentation(path)

    width, height = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    assert (round(width, 3), round(height, 3)) == (LayoutEngine.SLIDE_WIDTH, LayoutEngine.SLIDE_HEIGHT)
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert Emu(shape.left + shape.width).inches <= width + 1e-3
            assert Emu(shape.top + shape.height).inches <= height + 1e-3


def test_long_content_is_paginated_not_dropped(tmp_path):
    path = PresentationRenderer("Professional").render(_heavy_spec(), str(tmp_path / "deck.pptx"))
    prs = PPTXPresentation(path)
    texts = _texts(prs)

    # 9 spec slides, but long bullets/columns/table/summary continue on extra slides
    assert len(prs.slides) > 9
    assert any(t.endswith("(cont.)") for t in texts)
    assert any("(Part 1 of" in t for t in texts)
    for i in range(14):
        assert any(t.startswith(f"• {i}. ") for t in texts), f"bullet {i} missing"


def test_mismatched_chart_data_renders_and_empty_chart_is_not_invented(tmp_path):
    path = PresentationRenderer("Professional").render(_heavy_spec(), str(tmp_path / "deck.pptx"))
    prs = PPTXPresentation(path)
    charts = [sh.chart for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert len(charts) == 1
    assert list(charts[0].plots[0].categories) == ["Q1", "Q2"]
    assert any("No chart data" in t for t in _texts(prs))


def test_all_text_uses_theme_colours_on_dark_theme(tmp_path):
    """Regression: two-column and quote-author text used PowerPoint's default black (invisible on Dark)."""
    path = PresentationRenderer("Dark").render(_heavy_spec(), str(tmp_path / "deck.pptx"))
    prs = PPTXPresentation(path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    if run.text.strip():
                        assert run.font.color.type is not None, f"no explicit colour: {run.text[:40]!r}"
