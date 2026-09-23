from pptx import Presentation
from pptx.util import Inches

from app.document_processing.chunker import DocumentChunker
from app.document_processing.extractors.pptx_extractor import PPTXProcessor


def _make_pptx(path):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    s1.shapes.title.text = "Q3 Highlights"
    body = s1.placeholders[1].text_frame
    body.text = "Revenue grew 12%"
    body.add_paragraph().text = "Churn fell to 3.1%"
    s1.notes_slide.notes_text_frame.text = "Mention the EMEA launch."

    s2 = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    s2.shapes.title.text = "Regional Sales"
    table = s2.shapes.add_table(3, 2, Inches(1), Inches(1.5), Inches(6), Inches(2)).table
    for r, row in enumerate([["Region", "Sales"], ["EMEA", "1.9"], ["APAC", "1.1"]]):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    prs.save(path)


def test_pptx_extractor_reads_slides_tables_and_notes(tmp_path):
    path = str(tmp_path / "deck.pptx")
    _make_pptx(path)

    blocks = PPTXProcessor().process(path)

    text = next(b for b in blocks if not b["is_table"])
    assert text["page"] == 1 and text["section"] == "Q3 Highlights"
    assert "Revenue grew 12%" in text["content"] and "Churn fell to 3.1%" in text["content"]
    assert "Speaker notes: Mention the EMEA launch." in text["content"]

    table = next(b for b in blocks if b["is_table"])
    assert table["page"] == 2 and table["section"] == "Regional Sales"
    assert table["table_data"] == {"headers": ["Region", "Sales"], "rows": [["EMEA", "1.9"], ["APAC", "1.1"]]}


def test_chunker_routes_pptx_to_the_extractor(tmp_path):
    path = str(tmp_path / "deck.pptx")
    _make_pptx(path)

    chunks = DocumentChunker().process_and_chunk("doc-1", "deck.pptx", path)

    assert any(c["metadata"]["is_table"] for c in chunks)
    assert all(c["metadata"]["document_type"] == "pptx" for c in chunks)
    # Previously .pptx was read as plain text (zip bytes); now the content is real slide text
    assert any("Revenue grew 12%" in c["content"] for c in chunks)
