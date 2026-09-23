import json
from types import SimpleNamespace

import pytest
from pptx import Presentation as PPTXPresentation

from app.models import Project, Presentation, PresentationSlide, PresentationStatus
from app.rag import graph
from app.rag.graph import rag_engine

CONTEXTS = [
    {"content": "Revenue grew 12% year over year to $4.2M in Q3. Growth was driven by the EMEA region.",
     "document": "Q3_report.pdf", "page": 2, "section": "Revenue", "relevance_score": 0.9,
     "chunk_index": 0, "is_table": False, "table_data": None},
]


def _deck(env, slides):
    db, storage = env["db"], env["storage"]
    project = Project(user_id=env["owner"].id, name="Q3")
    db.add(project)
    db.commit()
    pres = Presentation(project_id=project.id, title="Q3 Deck", prompt="Q3 performance", theme="Professional",
                        status=PresentationStatus.COMPLETED)
    db.add(pres)
    db.commit()
    pres.pptx_path = storage.get_presentation_path(project.id, pres.id)
    for i, content in enumerate(slides, start=1):
        db.add(PresentationSlide(presentation_id=pres.id, slide_number=i, slide_type=content["type"],
                                 content_json=content, citations_json=[]))
    db.commit()
    return pres.id, pres.pptx_path


SLIDES = [
    {"type": "title", "title": "Q3 Deck", "subtitle": "Board", "citations": []},
    {"type": "bullet", "title": "Revenue", "bullets": ["old bullet"], "citations": []},
    {"type": "chart", "title": "Trend", "chart_type": "bar", "labels": ["Q1"], "values": [1.0], "citations": []},
]


@pytest.fixture
def no_llm(monkeypatch):
    monkeypatch.setattr(graph.retriever, "retrieve", lambda *a, **k: CONTEXTS)
    monkeypatch.setattr(rag_engine, "_get_llm", lambda: None)


def test_regenerate_keeps_type_and_rerenders_pptx(api_env, no_llm):
    pres_id, path = _deck(api_env, SLIDES)

    res = api_env["client"].post(f"/presentations/{pres_id}/slides/2/regenerate", json={"instructions": "focus on EMEA"})

    assert res.status_code == 200, res.text
    body = res.json()
    assert body["slide_type"] == "bullet"  # the old code replaced it with the new deck's first (title) slide
    assert body["content_json"]["bullets"][0].startswith("Revenue grew 12%")
    assert body["citations_json"][0]["document_name"] == "Q3_report.pdf"

    texts = [p.text for s in PPTXPresentation(path).slides for sh in s.shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs]
    assert any("Revenue grew 12%" in t for t in texts)
    assert not any("old bullet" in t for t in texts)


def test_regenerate_uses_llm_output_of_the_same_type(api_env, monkeypatch):
    monkeypatch.setattr(graph.retriever, "retrieve", lambda *a, **k: CONTEXTS)
    llm_slide = {"type": "title", "title": "EMEA growth", "bullets": ["EMEA drove Q3 growth"],
                 "citations": [{"document_name": "Q3_report.pdf", "page": 2}, {"document_name": "invented.pdf"}]}
    fake_llm = SimpleNamespace(invoke=lambda prompt: SimpleNamespace(content="```json\n" + json.dumps(llm_slide) + "\n```"))
    monkeypatch.setattr(rag_engine, "_get_llm", lambda: fake_llm)
    pres_id, _ = _deck(api_env, SLIDES)

    res = api_env["client"].post(f"/presentations/{pres_id}/slides/2/regenerate", json={})

    assert res.status_code == 200, res.text
    content = res.json()["content_json"]
    assert content["type"] == "bullet"
    assert content["bullets"] == ["EMEA drove Q3 growth"]
    assert [c["document_name"] for c in content["citations"]] == ["Q3_report.pdf"]


def test_regenerate_chart_without_llm_is_refused_not_invented(api_env, no_llm):
    pres_id, _ = _deck(api_env, SLIDES)
    res = api_env["client"].post(f"/presentations/{pres_id}/slides/3/regenerate", json={})
    assert res.status_code == 400
    assert "AI model" in res.json()["detail"]


def test_regenerate_other_users_deck_is_404(api_env, no_llm):
    pres_id, _ = _deck(api_env, SLIDES)
    api_env["current"]["user"] = api_env["other"]
    res = api_env["client"].post(f"/presentations/{pres_id}/slides/2/regenerate", json={})
    assert res.status_code == 404
    api_env["db"].expire_all()
    slide = api_env["db"].query(PresentationSlide).filter_by(presentation_id=pres_id, slide_number=2).one()
    assert slide.content_json["bullets"] == ["old bullet"]


def test_regenerate_uses_the_decks_saved_language_and_audience(api_env, monkeypatch):
    monkeypatch.setattr(graph.retriever, "retrieve", lambda *a, **k: CONTEXTS)
    prompts = []
    llm_slide = {"type": "bullet", "title": "Umsatz", "bullets": ["Der Umsatz wuchs um 12%"], "sources": ["S1"]}

    def invoke(prompt):
        prompts.append(prompt)
        return SimpleNamespace(content=json.dumps(llm_slide))

    monkeypatch.setattr(rag_engine, "_get_llm", lambda: SimpleNamespace(invoke=invoke))
    pres_id, _ = _deck(api_env, SLIDES)
    pres = api_env["db"].get(Presentation, pres_id)
    pres.language, pres.audience = "German", "Investors"
    api_env["db"].commit()

    res = api_env["client"].post(f"/presentations/{pres_id}/slides/2/regenerate", json={})

    assert res.status_code == 200, res.text
    assert "in German" in prompts[0] and "AUDIENCE: Investors" in prompts[0]
    assert res.json()["content_json"]["bullets"] == ["Der Umsatz wuchs um 12%"]
