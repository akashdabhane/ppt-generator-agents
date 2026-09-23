"""End-to-end without any API keys (PRD §6 Resilience): upload → ingest → generate → download a grounded .pptx."""
import io

from pptx import Presentation as PPTXPresentation
from pptx.util import Emu

from app.api.v1 import presentations as presentations_api
from app.core.config import settings
from app.models import Project
from app.rag import graph
from app.rag.vector_store import MockInMemoryVectorStore
from app.workers import tasks

REPORT = (
    "Q3 Revenue Review\n\n"
    "Revenue grew 12% year over year to $4.2M in Q3. Growth was driven by the EMEA region. "
    "Churn fell to 3.1% after the loyalty programme launched in July.\n\n"
    "Operating costs rose 4% because of new hires in the support team."
)
SALES_CSV = "Region,Sales\nEMEA,1.9\nAPAC,1.1\nAMER,1.2\n"


def test_upload_ingest_generate_download_without_api_keys(api_env, monkeypatch):
    for key in ("OPENAI_API_KEY", "ANTHROPIC_API_KEY", "GOOGLE_API_KEY"):
        monkeypatch.setattr(settings, key, None, raising=False)
    store = MockInMemoryVectorStore()
    monkeypatch.setattr(tasks, "SessionLocal", api_env["Session"])
    monkeypatch.setattr(tasks, "vector_store", store)
    monkeypatch.setattr(graph.retriever, "vector_store", store)
    monkeypatch.setattr(tasks, "storage_service", api_env["storage"])
    monkeypatch.setattr(presentations_api, "dispatch_presentation_generation",
                        lambda job_id, bg=None, **kw: tasks.run_presentation_generation(job_id, **kw))

    client, db = api_env["client"], api_env["db"]
    project = Project(user_id=api_env["owner"].id, name="Q3")
    db.add(project)
    db.commit()

    # Upload + ingest
    for name, body in [("q3_report.txt", REPORT), ("sales.csv", SALES_CSV)]:
        res = client.post(f"/projects/{project.id}/documents", files={"file": (name, io.BytesIO(body.encode()))})
        assert res.status_code == 200, res.text
    for doc_id in api_env["dispatched"]:
        tasks.run_document_ingestion(doc_id)
    docs = client.get(f"/projects/{project.id}/documents").json()
    assert [d["status"] for d in docs] == ["INDEXED", "INDEXED"], docs

    # Generate
    res = client.post(f"/projects/{project.id}/presentations/generate",
                      json={"prompt": "Q3 revenue and regional sales", "num_slides": 6, "theme": "Dark"})
    assert res.status_code == 200, res.text
    pres_id = res.json()["presentation_id"]
    progress = client.get(f"/presentations/{pres_id}/progress").json()
    assert progress["status"] == "COMPLETED", progress

    # Every content slide is cited, and only to the uploaded documents
    pres = client.get(f"/presentations/{pres_id}").json()
    content_slides = [s for s in pres["slides"] if s["slide_type"] != "title"]
    assert content_slides
    for s in content_slides:
        assert s["citations_json"], s
        assert {c["document_name"] for c in s["citations_json"]} <= {"q3_report.txt", "sales.csv"}
    assert any(s["slide_type"] == "table" for s in pres["slides"])
    assert "content slides cited" in pres["generation_summary"]

    # Download a valid 16:9 deck whose shapes all stay on the slide
    res = client.get(f"/presentations/{pres_id}/download")
    assert res.status_code == 200
    prs = PPTXPresentation(io.BytesIO(res.content))
    w, h = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    assert (round(w, 3), round(h, 3)) == (10.0, 5.625)
    for slide in prs.slides:
        for shape in slide.shapes:
            assert Emu(shape.left + shape.width).inches <= w + 1e-3
            assert Emu(shape.top + shape.height).inches <= h + 1e-3
