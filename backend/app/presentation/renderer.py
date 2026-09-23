import os
from typing import List, Dict, Any, Tuple
from pptx import Presentation as PPTXPresentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData

from app.presentation.themes import get_theme, PresentationTheme
from app.presentation.layout_engine import LayoutEngine, Rect
from app.schemas.presentation_spec import PresentationSpec, SlideSpec, Citation


class PresentationRenderer:
    BULLET_PREFIX = "• "
    TAKEAWAY_PREFIX = "✔  "
    BULLET_SPACE_AFTER = 12
    SUMMARY_FONT = 14
    SUMMARY_SPACE_AFTER = 14
    COLUMN_HEADING_FONT = 16
    COLUMN_BODY_FONT = 13
    COLUMN_SPACE_AFTER = 8
    QUOTE_FONT = 20
    AUTHOR_FONT = 14
    AUTHOR_SPACE_BEFORE = 12
    FOOTER_FONT = 9

    def __init__(self, theme_name: str = "Professional"):
        self.theme: PresentationTheme = get_theme(theme_name)

    def render(self, spec: PresentationSpec, output_path: str) -> str:
        prs = PPTXPresentation()

        # Set 16:9 Widescreen dimensions (10.0 x 5.625 inches)
        prs.slide_width = Inches(LayoutEngine.SLIDE_WIDTH)
        prs.slide_height = Inches(LayoutEngine.SLIDE_HEIGHT)

        # Use blank slide layout (index 6 in default template)
        blank_slide_layout = prs.slide_layouts[6]

        for slide_spec in spec.slides:
            s_dict = slide_spec.model_dump() if hasattr(slide_spec, "model_dump") else slide_spec
            s_type = s_dict.get("type", "bullet")
            title = s_dict.get("title", "")
            citations = s_dict.get("citations", [])

            if s_type == "table":
                # Handle Table Pagination across slides if necessary
                paginated_specs = LayoutEngine.split_table_rows(s_dict.get("columns", []), s_dict.get("rows", []))
                for p_spec in paginated_specs:
                    part_title = title or "Table Analysis"
                    if p_spec.get("total_parts", 1) > 1:
                        part_title += f" (Part {p_spec['part']} of {p_spec['total_parts']})"
                    slide = self._new_slide(prs, blank_slide_layout)
                    self._render_header(slide, part_title)
                    self._render_table(slide, p_spec["columns"], p_spec["rows"])
                    self._render_footer(slide, citations)
            elif s_type in ("title", "section"):
                slide = self._new_slide(prs, blank_slide_layout)
                self._render_title_slide(slide, title, s_dict.get("subtitle") or "")
            elif s_type in ("bullet", "summary", "two_column"):
                # Text slides paginate like tables: long lists continue on "(cont.)" slides
                if s_type == "bullet":
                    pages = self._paginate_bullets(s_dict.get("bullets", []))
                elif s_type == "summary":
                    pages = self._paginate_takeaways(s_dict.get("key_takeaways", []))
                else:
                    pages = self._paginate_two_column(s_dict)
                for idx, page in enumerate(pages):
                    slide = self._new_slide(prs, blank_slide_layout)
                    self._render_header(slide, title if idx == 0 else f"{title} (cont.)")
                    if s_type == "bullet":
                        self._render_bullet_slide(slide, page)
                    elif s_type == "summary":
                        self._render_summary_slide(slide, page)
                    else:
                        self._render_two_column_slide(slide, s_dict, *page)
                    self._render_footer(slide, citations)
            elif s_type == "chart":
                slide = self._new_slide(prs, blank_slide_layout)
                self._render_header(slide, title)
                self._render_chart_slide(slide, s_dict)
                self._render_footer(slide, citations)
            elif s_type == "quote":
                slide = self._new_slide(prs, blank_slide_layout)
                self._render_header(slide, title)
                self._render_quote_slide(slide, s_dict)
                self._render_footer(slide, citations)
            else:
                slide = self._new_slide(prs, blank_slide_layout)
                self._render_header(slide, title or "Slide")
                self._render_bullet_slide(slide, ["Content block"])
                self._render_footer(slide, citations)

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        prs.save(output_path)
        return output_path

    # --- Helpers ---------------------------------------------------------------

    def _new_slide(self, prs, layout):
        slide = prs.slides.add_slide(layout)
        self._apply_background(slide)
        return slide

    def _apply_background(self, slide):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self._rgb(self.theme.colors.background)

    @staticmethod
    def _rgb(color) -> RGBColor:
        return RGBColor(*color)

    def _write(self, p, text: str, size: float, color, font: str = None, bold: bool = False,
               italic: bool = False, align=None):
        """Sets paragraph text with formatting on the run itself (not just the paragraph default),
        which PowerPoint, Google Slides, Keynote and LibreOffice all honour."""
        p.text = text
        if align is not None:
            p.alignment = align
        for f in [p.font] + [r.font for r in p.runs]:
            f.name = font or self.theme.secondary_font
            f.size = Pt(size)
            f.bold = bold
            f.italic = italic
            f.color.rgb = self._rgb(color)

    @staticmethod
    def _add_textbox(slide, rect: Rect):
        box = slide.shapes.add_textbox(Inches(rect.x), Inches(rect.y), Inches(rect.width), Inches(rect.height))
        tf = box.text_frame
        tf.word_wrap = True
        return tf

    # --- Pagination (all measurement lives in LayoutEngine) --------------------

    def _paginate_bullets(self, bullets: List[str]) -> List[List[str]]:
        w, h = LayoutEngine.usable_size(LayoutEngine.get_content_rect())
        return LayoutEngine.paginate_text_items(
            bullets, w, h, self.theme.body_size, self.BULLET_SPACE_AFTER, self.BULLET_PREFIX)

    def _paginate_takeaways(self, takeaways: List[str]) -> List[List[str]]:
        w, h = LayoutEngine.usable_size(LayoutEngine.get_content_rect())
        return LayoutEngine.paginate_text_items(
            takeaways, w, h, self.SUMMARY_FONT, self.SUMMARY_SPACE_AFTER, self.TAKEAWAY_PREFIX, bold=True)

    def _column_heading(self, text: str) -> Tuple[str, float]:
        """Column heading limited to two lines; returns the text and its height."""
        w, _ = LayoutEngine.usable_size(LayoutEngine.get_two_column_rects()[0])
        text = LayoutEngine.truncate_to_fit(
            text, w, 2 * LayoutEngine.line_height(self.COLUMN_HEADING_FONT), self.COLUMN_HEADING_FONT, bold=True)
        return text, LayoutEngine.text_height(text, w, self.COLUMN_HEADING_FONT, bold=True)

    def _paginate_two_column(self, s_dict: Dict[str, Any]) -> List[Tuple[List[str], List[str]]]:
        w, h = LayoutEngine.usable_size(LayoutEngine.get_two_column_rects()[0])
        columns = []
        for side in ("left", "right"):
            _, heading_h = self._column_heading(s_dict.get(f"{side}_title", ""))
            columns.append(LayoutEngine.paginate_text_items(
                s_dict.get(f"{side}_content", []), w, h - heading_h, self.COLUMN_BODY_FONT,
                self.COLUMN_SPACE_AFTER, self.BULLET_PREFIX))
        left, right = columns
        n = max(len(left), len(right))
        return [(left[i] if i < len(left) else [], right[i] if i < len(right) else []) for i in range(n)]

    # --- Drawing ---------------------------------------------------------------

    def _render_header(self, slide, title_text: str):
        rect = LayoutEngine.get_title_rect()
        w, h = LayoutEngine.usable_size(rect)
        size, text = LayoutEngine.fit_font_size(title_text, w, h, self.theme.heading_size, bold=True)
        self._write(self._add_textbox(slide, rect).paragraphs[0], text, size, self.theme.colors.primary,
                    self.theme.primary_font, bold=True)

    def _render_footer(self, slide, citations: List[Dict[str, Any]]):
        if not citations:
            return
        rect = LayoutEngine.get_footer_rect()
        w, h = LayoutEngine.usable_size(rect)
        c_text = "Sources: " + ", ".join(
            [f"{c.get('document_name', 'Doc')} (p. {c.get('page') or 'N/A'})" for c in citations[:2]])
        self._write(self._add_textbox(slide, rect).paragraphs[0], LayoutEngine.truncate_to_fit(c_text, w, h, self.FOOTER_FONT),
                    self.FOOTER_FONT, self.theme.colors.secondary, italic=True)

    def _render_title_slide(self, slide, title: str, subtitle: str):
        rect = LayoutEngine.get_title_card_rect()
        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(rect.x), Inches(rect.y), Inches(rect.width), Inches(rect.height)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = self._rgb(self.theme.colors.card_bg)
        card.line.color.rgb = self._rgb(self.theme.colors.card_bg)

        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        w, h = LayoutEngine.usable_size(rect)

        # The title gets up to 60 % of the card (all of it without a subtitle); the subtitle gets the rest
        title_budget = h * 0.6 if subtitle else h
        title_size, title = LayoutEngine.fit_font_size(title, w, title_budget, self.theme.title_size, 20, bold=True)
        self._write(tf.paragraphs[0], title, title_size, self.theme.colors.primary, self.theme.primary_font,
                    bold=True, align=PP_ALIGN.CENTER)

        if subtitle:
            used = LayoutEngine.text_height(title, w, title_size, bold=True)
            sub_size, subtitle = LayoutEngine.fit_font_size(subtitle, w, h - used, self.theme.subtitle_size)
            self._write(tf.add_paragraph(), subtitle, sub_size, self.theme.colors.secondary, align=PP_ALIGN.CENTER)

    def _render_bullet_slide(self, slide, bullets: List[str]):
        tf = self._add_textbox(slide, LayoutEngine.get_content_rect())
        for idx, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            self._write(p, f"{self.BULLET_PREFIX}{bullet}", self.theme.body_size, self.theme.colors.text)
            p.space_after = Pt(self.BULLET_SPACE_AFTER)

    def _render_column(self, slide, rect: Rect, heading: str, items: List[str]):
        tf = self._add_textbox(slide, rect)
        self._write(tf.paragraphs[0], self._column_heading(heading)[0], self.COLUMN_HEADING_FONT,
                    self.theme.colors.primary, self.theme.primary_font, bold=True)

        for item in items:
            p = tf.add_paragraph()
            self._write(p, f"{self.BULLET_PREFIX}{item}", self.COLUMN_BODY_FONT, self.theme.colors.text)
            p.space_after = Pt(self.COLUMN_SPACE_AFTER)

    def _render_two_column_slide(self, slide, s_dict: Dict[str, Any], left_items: List[str], right_items: List[str]):
        left_rect, right_rect = LayoutEngine.get_two_column_rects()
        self._render_column(slide, left_rect, s_dict.get("left_title", "Left Column"), left_items)
        self._render_column(slide, right_rect, s_dict.get("right_title", "Right Column"), right_items)

    def _render_table(self, slide, columns: List[str], rows: List[List[str]]):
        rect = LayoutEngine.get_content_rect()
        num_rows = len(rows) + 1
        num_cols = max(len(columns), 1)

        table_shape = slide.shapes.add_table(
            num_rows, num_cols,
            Inches(rect.x), Inches(rect.y), Inches(rect.width), Inches(rect.height)
        )
        table = table_shape.table

        # Populate headers
        for col_idx, col_name in enumerate(columns):
            cell = table.cell(0, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = self._rgb(self.theme.colors.primary)
            self._write(cell.text_frame.paragraphs[0], str(col_name), LayoutEngine.TABLE_HEADER_FONT_PT,
                        (255, 255, 255), bold=True)

        # Populate rows
        for row_idx, row in enumerate(rows, start=1):
            bg_color = self.theme.colors.card_bg if row_idx % 2 == 0 else self.theme.colors.background
            for col_idx, val in enumerate(row):
                if col_idx < num_cols:
                    cell = table.cell(row_idx, col_idx)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self._rgb(bg_color)
                    self._write(cell.text_frame.paragraphs[0], str(val), LayoutEngine.TABLE_BODY_FONT_PT,
                                self.theme.colors.text)

    def _render_chart_slide(self, slide, s_dict: Dict[str, Any]):
        rect = LayoutEngine.get_content_rect()
        chart_type_str = s_dict.get("chart_type", "bar")
        labels = s_dict.get("labels", [])
        values = s_dict.get("values", [])
        n = min(len(labels), len(values))
        if n == 0:
            # Never draw a chart with invented numbers
            self._render_bullet_slide(slide, ["No chart data was found in the source documents."])
            return

        chart_data = CategoryChartData()
        chart_data.categories = [str(label) for label in labels[:n]]
        chart_data.add_series("Series 1", values[:n])

        xl_chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED
        if chart_type_str == "line":
            xl_chart_type = XL_CHART_TYPE.LINE
        elif chart_type_str == "pie":
            xl_chart_type = XL_CHART_TYPE.PIE

        slide.shapes.add_chart(
            xl_chart_type,
            Inches(rect.x), Inches(rect.y), Inches(rect.width), Inches(rect.height),
            chart_data
        )

    def _render_quote_slide(self, slide, s_dict: Dict[str, Any]):
        rect = LayoutEngine.get_content_rect()
        w, h = LayoutEngine.usable_size(rect)
        tf = self._add_textbox(slide, rect)

        author = s_dict.get("author")
        author_h = 0.0
        if author:
            author = LayoutEngine.truncate_to_fit(
                f"— {author}", w, 2 * LayoutEngine.line_height(self.AUTHOR_FONT), self.AUTHOR_FONT, bold=True)
            author_h = LayoutEngine.text_height(author, w, self.AUTHOR_FONT, bold=True) + self.AUTHOR_SPACE_BEFORE / 72.0

        size, quote = LayoutEngine.fit_font_size(f"“{s_dict.get('quote', '')}”", w, h - author_h, self.QUOTE_FONT)
        self._write(tf.paragraphs[0], quote, size, self.theme.colors.primary, italic=True)

        if author:
            p2 = tf.add_paragraph()
            self._write(p2, author, self.AUTHOR_FONT, self.theme.colors.text, bold=True)
            p2.space_before = Pt(self.AUTHOR_SPACE_BEFORE)

    def _render_summary_slide(self, slide, takeaways: List[str]):
        tf = self._add_textbox(slide, LayoutEngine.get_content_rect())
        for idx, item in enumerate(takeaways):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            self._write(p, f"{self.TAKEAWAY_PREFIX}{item}", self.SUMMARY_FONT, self.theme.colors.accent, bold=True)
            p.space_after = Pt(self.SUMMARY_SPACE_AFTER)
