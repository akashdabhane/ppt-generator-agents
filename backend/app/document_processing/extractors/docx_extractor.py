import docx
from typing import List, Dict, Any


class DOCXProcessor:
    def process(self, file_path: str) -> List[Dict[str, Any]]:
        doc = docx.Document(file_path)
        blocks = []
        current_heading = "Document Content"

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style.name.startswith("Heading"):
                current_heading = text
                continue
            
            blocks.append({
                "content": text,
                "page": None,
                "section": current_heading,
                "is_table": False,
                "table_data": None
            })

        for table_idx, table in enumerate(doc.tables, start=1):
            if not table.rows:
                continue
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows[1:]]

            table_str = f"Table {table_idx}:\n"
            table_str += " | ".join(headers) + "\n"
            table_str += "-" * 40 + "\n"
            for row in rows:
                table_str += " | ".join(row) + "\n"

            blocks.append({
                "content": table_str.strip(),
                "page": None,
                "section": f"Table {table_idx}",
                "is_table": True,
                "table_data": {
                    "headers": headers,
                    "rows": rows
                }
            })

        return blocks
