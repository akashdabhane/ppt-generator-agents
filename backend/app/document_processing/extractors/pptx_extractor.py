from typing import List, Dict, Any, Iterable
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


class PPTXProcessor:
    """Extracts per-slide text (page = slide number, section = slide title), tables and speaker notes."""

    def process(self, file_path: str) -> List[Dict[str, Any]]:
        prs = Presentation(file_path)
        blocks = []

        for slide_no, slide in enumerate(prs.slides, start=1):
            title_shape = slide.shapes.title
            title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
            section = title or f"Slide {slide_no}"

            texts = []
            for shape in self._walk(slide.shapes):
                if shape is title_shape:
                    continue
                if shape.has_table:
                    table_block = self._table_block(shape.table, slide_no, section)
                    if table_block:
                        blocks.append(table_block)
                elif shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        line = "".join(r.text for r in p.runs).strip()
                        if line:
                            texts.append(line)

            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
                if notes:
                    texts.append(f"Speaker notes: {notes}")

            if texts:
                blocks.append({
                    "content": "\n".join(texts),
                    "page": slide_no,
                    "section": section,
                    "is_table": False,
                    "table_data": None
                })

        return blocks

    def _walk(self, shapes) -> Iterable:
        """Yields shapes, descending into groups."""
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from self._walk(shape.shapes)
            else:
                yield shape

    @staticmethod
    def _table_block(table, slide_no: int, section: str):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows:
            return None
        headers, body = rows[0], rows[1:]

        table_str = f"Table on slide {slide_no} ({section}):\n"
        table_str += " | ".join(headers) + "\n"
        table_str += "-" * 40 + "\n"
        for row in body:
            table_str += " | ".join(row) + "\n"

        return {
            "content": table_str.strip(),
            "page": slide_no,
            "section": section,
            "is_table": True,
            "table_data": {
                "headers": headers,
                "rows": body
            }
        }
